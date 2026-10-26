"""Shared python-docx layout for the exported files (plan §8.2, §8.3): A4 page, serif body, the draft banner and label.

Every exported DOCX carries the banner "AI-generated draft: check against your court's rules." in its header and the
brief's label (`DRAFT_LABEL`) in its footer, so neither can be lost when one page is printed on its own (plan §8.1, §16).
python-docx is imported inside the functions, like every heavy dependency of the package.
"""

from __future__ import annotations

import re
from typing import Any

from vakeel_sahayak.schemas.bail_brief import DRAFT_LABEL

BANNER = "AI-generated draft: check against your court's rules."
FONT = "Times New Roman"
BODY_PT = 12
_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")


def court_date(iso: str | None) -> str:
    """'2024-08-20' as '20.08.2024', the way Indian court filings write dates; anything else is returned unchanged."""
    m = _ISO.match(iso or "")
    return f"{m.group(3)}.{m.group(2)}.{m.group(1)}" if m else (iso or "")


def new_document(title: str | None = None) -> Any:
    """An empty A4 document with the body font, margins, the banner in the header and the label and page number in the footer.

    Args:
        title: document title stored in the file's core properties

    Returns:
        A `docx.Document`.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), FONT)  # complex scripts (Devanagari quotes)
    normal.paragraph_format.space_after = Pt(6)
    if title:
        doc.core_properties.title = title

    banner = sec.header.paragraphs[0]
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = banner.add_run(BANNER)
    run.bold, run.italic, run.font.size = True, True, Pt(9)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = footer.add_run(f"{DRAFT_LABEL}  ·  Page ")
    label.font.size = Pt(9)
    _page_field(footer.add_run())
    return doc


def _page_field(run: Any) -> None:
    """A PAGE field in `run` (Word fills in the number; the file itself holds no page text)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.size = Pt(9)
    for kind, instr in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        run._r.append(el)


def heading(doc: Any, text: str, *, size: int = BODY_PT, space_before: int = 12) -> Any:
    """A centred, bold section heading kept with the paragraph that follows it."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold, run.font.size = True, Pt(size)
    return p


def paragraph(doc: Any, text: str = "", *, bold: bool = False, italic: bool = False, align: str | None = None,
              indent_cm: float | None = None, hanging_cm: float | None = None, justify: bool = True) -> Any:
    """A body paragraph; `align` is 'center' or 'right' to override the default justified text."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    p = doc.add_paragraph()
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(
        align or "", WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    )
    if indent_cm is not None:
        p.paragraph_format.left_indent = Cm(indent_cm)
    if hanging_cm is not None:
        p.paragraph_format.first_line_indent = Cm(-hanging_cm)
    if text:
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
    return p


def table(doc: Any, header: list[str], rows: list[list[str]], widths_cm: list[float] | None = None) -> Any:
    """A bordered table with a bold header row that repeats on every page."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm

    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, header, strict=True):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            cell.text = text
    if widths_cm:
        for row in t.rows:
            for cell, w in zip(row.cells, widths_cm, strict=True):
                cell.width = Cm(w)
    return t
